"""
Generate the Render vs Trigger.dev cost comparison Excel model.
Run: python3 cost_model.py
Output: COST_MODEL.xlsx
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Styles ──────────────────────────────────────────────────────────────────

header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
section_font = Font(bold=True, size=12, color="2F5496")
input_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
result_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
warning_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
label_font = Font(bold=True, size=10)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
money_fmt = '$#,##0.00'
money_fmt_4 = '$#,##0.0000'
pct_fmt = '0.0%'
int_fmt = '#,##0'
rate_fmt = '$0.0000000'


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border


def style_input(ws, row, col):
    cell = ws.cell(row=row, column=col)
    cell.fill = input_fill
    cell.border = thin_border
    return cell


def style_result(ws, row, col):
    cell = ws.cell(row=row, column=col)
    cell.fill = result_fill
    cell.border = thin_border
    cell.font = Font(bold=True)
    return cell


def style_cell(ws, row, col):
    cell = ws.cell(row=row, column=col)
    cell.border = thin_border
    return cell


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 1: RATES
# ═══════════════════════════════════════════════════════════════════════════

ws1 = wb.active
ws1.title = "Rates"
ws1.sheet_properties.tabColor = "2F5496"

# Title
ws1["A1"] = "Compute Rates Reference"
ws1["A1"].font = Font(bold=True, size=14, color="2F5496")
ws1.merge_cells("A1:F1")

# Render rates
r = 3
ws1.cell(r, 1, "Render Workflows").font = section_font
r = 4
for c, h in enumerate(["Tier", "CPU", "Memory", "$/hour", "$/second", "$/second (formula)"], 1):
    ws1.cell(r, c, h)
style_header(ws1, r, 6)

render_tiers = [
    ("Starter", 0.5, "512 MB", 0.05),
    ("Standard", 1, "2 GB", 0.20),
    ("Pro", 2, "4 GB", 0.40),
    ("Pro Plus", 4, "8 GB", 1.00),
    ("Pro Max", 4, "16 GB", 2.00),
    ("Pro Ultra", 8, "32 GB", 7.00),
]
for i, (name, cpu, mem, hourly) in enumerate(render_tiers):
    row = r + 1 + i
    style_cell(ws1, row, 1).value = name
    style_cell(ws1, row, 2).value = cpu
    style_cell(ws1, row, 3).value = mem
    style_cell(ws1, row, 4).value = hourly
    ws1.cell(row, 4).number_format = money_fmt
    style_cell(ws1, row, 5).value = hourly / 3600
    ws1.cell(row, 5).number_format = rate_fmt
    style_cell(ws1, row, 6).value = f"=D{row}/3600"
    ws1.cell(row, 6).number_format = rate_fmt

# Trigger rates
r = r + len(render_tiers) + 2
ws1.cell(r, 1, "Trigger.dev").font = section_font
r += 1
for c, h in enumerate(["Tier", "CPU", "Memory", "$/second", "Invocation $/run", ""], 1):
    ws1.cell(r, c, h)
style_header(ws1, r, 5)

trigger_tiers = [
    ("Micro", 0.25, "0.25 GB", 0.0000169),
    ("Small 1x", 0.5, "0.5 GB", 0.0000338),
    ("Small 2x", 1, "1 GB", 0.0000675),
    ("Medium 1x", 1, "2 GB", 0.0000850),
    ("Medium 2x", 2, "4 GB", 0.0001700),
    ("Large 1x", 4, "8 GB", 0.0003400),
    ("Large 2x", 8, "16 GB", 0.0006800),
]
for i, (name, cpu, mem, per_s) in enumerate(trigger_tiers):
    row = r + 1 + i
    style_cell(ws1, row, 1).value = name
    style_cell(ws1, row, 2).value = cpu
    style_cell(ws1, row, 3).value = mem
    style_cell(ws1, row, 4).value = per_s
    ws1.cell(row, 4).number_format = rate_fmt
    style_cell(ws1, row, 5).value = 0.000025
    ws1.cell(row, 5).number_format = rate_fmt

# BG Worker rates
r = r + len(trigger_tiers) + 2
ws1.cell(r, 1, "Background Workers (Render)").font = section_font
r += 1
for c, h in enumerate(["Tier", "CPU", "Memory", "$/month", "$/hour", "$/second (= $/month ÷ 730hrs ÷ 3600s)"], 1):
    ws1.cell(r, c, h)
style_header(ws1, r, 6)

bg_tiers = [
    ("Starter", 0.5, "512 MB", 7),
    ("Standard", 1, "2 GB", 25),
    ("Pro", 2, "4 GB", 85),
    ("Pro Plus", 4, "8 GB", 175),
    ("Pro Max", 4, "16 GB", 225),
    ("Pro Ultra", 8, "32 GB", 450),
]
for i, (name, cpu, mem, monthly) in enumerate(bg_tiers):
    row = r + 1 + i
    style_cell(ws1, row, 1).value = name
    style_cell(ws1, row, 2).value = cpu
    style_cell(ws1, row, 3).value = mem
    style_cell(ws1, row, 4).value = monthly
    ws1.cell(row, 4).number_format = money_fmt
    # $/hour = monthly / 730 hours
    style_cell(ws1, row, 5).value = f"=D{row}/730"
    ws1.cell(row, 5).number_format = '$#,##0.0000'
    # $/second = monthly / 730 / 3600
    style_cell(ws1, row, 6).value = f"=D{row}/730/3600"
    ws1.cell(row, 6).number_format = rate_fmt

ws1.column_dimensions["A"].width = 18
ws1.column_dimensions["B"].width = 10
ws1.column_dimensions["C"].width = 12
ws1.column_dimensions["D"].width = 14
ws1.column_dimensions["E"].width = 16
ws1.column_dimensions["F"].width = 18


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 2: TASK TYPE COMPARISON
# ═══════════════════════════════════════════════════════════════════════════

ws2 = wb.create_sheet("By Task Type")
ws2.sheet_properties.tabColor = "548235"

ws2["A1"] = "Cost by Task Type — Workflows Today vs After Each Project"
ws2["A1"].font = Font(bold=True, size=14, color="2F5496")
ws2.merge_cells("A1:L1")

ws2["A2"] = "Yellow cells are inputs. Shows Workflows cost at each project stage vs Trigger.dev and BG Worker."
ws2["A2"].font = Font(italic=True, size=9, color="666666")

# Rates
r = 4
ws2.cell(r, 1, "Workflows $/second").font = label_font
style_input(ws2, r, 2).value = 0.0000556
ws2.cell(r, 2).number_format = rate_fmt
ws2.cell(r, 3, "(Standard: 1 CPU, 2 GB)").font = Font(size=9, italic=True, color="888888")
r = 5
ws2.cell(r, 1, "Trigger.dev $/second").font = label_font
style_input(ws2, r, 2).value = 0.0000850
ws2.cell(r, 2).number_format = rate_fmt
ws2.cell(r, 3, "(Medium 1x — has active compute + waitpoints)").font = Font(size=9, italic=True, color="888888")
r = 6
ws2.cell(r, 1, "Trigger.dev invocation $/run").font = label_font
style_input(ws2, r, 2).value = 0.000025
ws2.cell(r, 2).number_format = rate_fmt
r = 7
ws2.cell(r, 1, "BG Worker $/month").font = label_font
style_input(ws2, r, 2).value = 85
ws2.cell(r, 2).number_format = money_fmt
ws2.cell(r, 3, "(Pro: 2 CPU, 4 GB — 1 instance, 2 subprocess workers)").font = Font(size=9, italic=True, color="888888")

rr = "$B$4"  # Render rate
tr = "$B$5"  # Trigger rate
ti = "$B$6"  # Trigger invocation
bg = "$B$7"  # BG Worker monthly

# Header
r = 9
headers = [
    "Task Type", "Monthly\nVolume",
    "Wall-Clock\n(s)", "Active\nCompute (s)", "Waitpoint\nCompute (s)", "Runs/\nReq",
    "Workflows\nTODAY",
    "After P1:\nNo Parent\nBilling",
    "After P2:\nActive\nCompute",
    "After P3:\nWaitpoints",
    "Trigger.dev", "BG Worker",
    "Workflows\nWins After"
]
for c, h in enumerate(headers, 1):
    ws2.cell(r, c, h)
style_header(ws2, r, len(headers))

tasks = [
    ("Simple compute",        500000,  1,    1,    1,    1),
    ("Image processing",      200000,  5,    5,    5,    1),
    ("Single LLM call",       300000,  8,    0.5,  0.3,  1),
    ("Multi-step LLM agent",  100000,  35,   2.5,  1.5,  6),
    ("AI pipeline",            50000,  90,   6,    3,    4),
    ("Data fan-out",           10000,  200,  90,   85,   121),
    ("Long-running agent",     10000,  350,  15,   5,    10),
]

for i, (name, vol, wc, ac, wp, runs) in enumerate(tasks):
    row = r + 1 + i
    v = f"B{row}"

    style_cell(ws2, row, 1).value = name
    ws2.cell(row, 1).font = label_font
    style_input(ws2, row, 2).value = vol
    ws2.cell(row, 2).number_format = int_fmt
    style_input(ws2, row, 3).value = wc
    style_input(ws2, row, 4).value = ac
    style_input(ws2, row, 5).value = wp
    style_input(ws2, row, 6).value = runs

    # G: Workflows TODAY = vol × wall_clock × rate
    style_result(ws2, row, 7)
    ws2.cell(row, 7).value = f"={v}*C{row}*{rr}"
    ws2.cell(row, 7).number_format = money_fmt

    # H: After P1 — parent active, subtasks wall-clock
    style_result(ws2, row, 8)
    ws2.cell(row, 8).value = (
        f"=IF(F{row}=1,G{row},"
        f"{v}*(D{row}+(C{row}-D{row})*(F{row}-1)/F{row})*{rr})"
    )
    ws2.cell(row, 8).number_format = money_fmt

    # I: After P2 — all tasks active compute only
    style_result(ws2, row, 9)
    ws2.cell(row, 9).value = f"={v}*D{row}*{rr}"
    ws2.cell(row, 9).number_format = money_fmt

    # J: After P3 — waitpoint compute
    style_result(ws2, row, 10)
    ws2.cell(row, 10).value = f"={v}*E{row}*{rr}"
    ws2.cell(row, 10).number_format = money_fmt

    # K: Trigger.dev (already has active compute + waitpoints)
    style_result(ws2, row, 11)
    ws2.cell(row, 11).value = f"={v}*E{row}*{tr}+{v}*F{row}*{ti}"
    ws2.cell(row, 11).number_format = money_fmt

    # L: BG Worker (flat)
    style_result(ws2, row, 12)
    ws2.cell(row, 12).value = f"={bg}"
    ws2.cell(row, 12).number_format = money_fmt

    # M: Which project makes Workflows win?
    style_result(ws2, row, 13)
    ws2.cell(row, 13).value = (
        f'=IF(AND(G{row}<=K{row},G{row}<=L{row}),"Already wins",'
        f'IF(AND(H{row}<=K{row},H{row}<=L{row}),"P1: No parent billing",'
        f'IF(AND(I{row}<=K{row},I{row}<=L{row}),"P2: Active compute",'
        f'IF(AND(J{row}<=K{row},J{row}<=L{row}),"P3: Waitpoints",'
        f'"Cannot win"))))'
    )
    ws2.cell(row, 13).font = Font(bold=True, size=10)

# Notes
nr = r + len(tasks) + 2
notes = [
    "P1 — No parent billing: Stop billing parent tasks while subtasks run. Billing change only.",
    "P2 — Active compute: Bill CPU-active time across all tasks. Requires cgroup CPU tracking.",
    "P3 — Waitpoints: Explicit SDK primitives to pause billing during LLM/API calls. Requires CRIU. Builds on P2.",
    "Trigger.dev already has P2 + P3. Their column uses Waitpoint Compute because their billing reflects both.",
    "BG Worker: flat $85/mo (1× Pro). Does NOT include retries, orchestration, observability, or burst handling.",
]
for i, note in enumerate(notes):
    ws2.cell(nr + i, 1, note).font = Font(size=9, italic=True, color="666666")
    ws2.merge_cells(start_row=nr + i, start_column=1, end_row=nr + i, end_column=10)

for c in range(1, 14):
    ws2.column_dimensions[get_column_letter(c)].width = [24, 12, 10, 10, 10, 7, 13, 13, 13, 13, 13, 11, 22][c - 1]


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 3: CUSTOMER SCENARIO
# ═══════════════════════════════════════════════════════════════════════════

ws3 = wb.create_sheet("Customer Scenario")
ws3.sheet_properties.tabColor = "BF8F00"

ws3["A1"] = "AI Support Customer — Cost Model"
ws3["A1"].font = Font(bold=True, size=14, color="2F5496")
ws3.merge_cells("A1:F1")
ws3["A2"] = "Change yellow cells to model your own workload."
ws3["A2"].font = Font(italic=True, size=9, color="666666")

# ── INPUTS ──
r = 4
ws3.cell(r, 1, "INPUTS").font = section_font
r = 5
inputs = [
    ("Requests per month", 800000, int_fmt),
    ("Subtasks per request", 10, int_fmt),
    ("Parent wall-clock (seconds)", 20, "0.0"),
    ("Parent active compute (seconds)", 2, "0.0"),
    ("Subtask wall-clock (seconds)", 3, "0.0"),
    ("Subtask active compute (seconds)", 0.5, "0.0"),
]
for i, (label, val, fmt) in enumerate(inputs):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    style_input(ws3, row, 2).value = val
    ws3.cell(row, 2).number_format = fmt

# Named references for readability
req_cell = "B5"     # requests/month
subs_cell = "B6"    # subtasks per request
p_wc_cell = "B7"    # parent wall-clock
p_ac_cell = "B8"    # parent active compute
s_wc_cell = "B9"    # subtask wall-clock
s_ac_cell = "B10"   # subtask active compute

r = 12
ws3.cell(r, 1, "RATES").font = section_font
r = 13
rate_inputs = [
    ("Render $/second (Standard)", 0.0000556, rate_fmt),
    ("Trigger $/second (Medium 1x)", 0.0000850, rate_fmt),
    ("Trigger invocation $/run", 0.000025, rate_fmt),
]
for i, (label, val, fmt) in enumerate(rate_inputs):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    style_input(ws3, row, 2).value = val
    ws3.cell(row, 2).number_format = fmt

r_rate = "B13"
t_rate = "B14"
t_invoc = "B15"

# ── DERIVED ──
r = 17
ws3.cell(r, 1, "DERIVED VALUES").font = section_font
r = 18
derived = [
    ("Total task runs / month", f"={req_cell}*(1+{subs_cell})", int_fmt),
    ("Total parent billed seconds (Render)", f"={req_cell}*{p_wc_cell}", int_fmt),
    ("Total subtask billed seconds (Render)", f"={req_cell}*{subs_cell}*{s_wc_cell}", int_fmt),
    ("Total parent billed seconds (Trigger)", f"={req_cell}*{p_ac_cell}", int_fmt),
    ("Total subtask billed seconds (Trigger)", f"={req_cell}*{subs_cell}*{s_ac_cell}", int_fmt),
    ("Wait ratio (parent)", f"=1-{p_ac_cell}/{p_wc_cell}", pct_fmt),
    ("Wait ratio (subtask)", f"=1-{s_ac_cell}/{s_wc_cell}", pct_fmt),
]
for i, (label, formula, fmt) in enumerate(derived):
    row = r + i
    style_cell(ws3, row, 1).value = label
    style_cell(ws3, row, 2).value = formula
    ws3.cell(row, 2).number_format = fmt

# ── RENDER WORKFLOWS (TODAY) ──
r = 27
ws3.cell(r, 1, "RENDER WORKFLOWS — TODAY").font = section_font
ws3.cell(r, 1).fill = warning_fill
r = 28
render_today = [
    ("Parent compute cost", f"={req_cell}*{p_wc_cell}*{r_rate}", money_fmt),
    ("Subtask compute cost", f"={req_cell}*{subs_cell}*{s_wc_cell}*{r_rate}", money_fmt),
    ("Invocation fees", 0, money_fmt),
    ("TOTAL", f"=B28+B29+B30", money_fmt),
]
for i, (label, formula, fmt) in enumerate(render_today):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    if i == len(render_today) - 1:
        style_result(ws3, row, 2).value = formula
    else:
        style_cell(ws3, row, 2).value = formula
    ws3.cell(row, 2).number_format = fmt

# ── APPROACH 1 ──
r = 34
ws3.cell(r, 1, "APPROACH 1 — No parent billing during subtask waits").font = section_font
ws3.cell(r, 1).fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
r = 35
approach1 = [
    ("Parent compute cost (active only)", f"={req_cell}*{p_ac_cell}*{r_rate}", money_fmt),
    ("Subtask compute cost (unchanged)", f"={req_cell}*{subs_cell}*{s_wc_cell}*{r_rate}", money_fmt),
    ("Invocation fees", 0, money_fmt),
    ("TOTAL", f"=B35+B36+B37", money_fmt),
    ("Savings vs today", f"=B31-B38", money_fmt),
    ("% reduction", f"=1-B38/B31", pct_fmt),
    ("vs Trigger.dev", f"=B38/B56", '0.0"x"'),
]
for i, (label, formula, fmt) in enumerate(approach1):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    if label == "TOTAL":
        style_result(ws3, row, 2).value = formula
    else:
        style_cell(ws3, row, 2).value = formula
    ws3.cell(row, 2).number_format = fmt

# ── APPROACH 2 ──
r = 43
ws3.cell(r, 1, "APPROACH 2 — Active compute billing (all tasks)").font = section_font
ws3.cell(r, 1).fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
r = 44
approach2 = [
    ("Parent compute cost (active only)", f"={req_cell}*{p_ac_cell}*{r_rate}", money_fmt),
    ("Subtask compute cost (active only)", f"={req_cell}*{subs_cell}*{s_ac_cell}*{r_rate}", money_fmt),
    ("Invocation fees", 0, money_fmt),
    ("TOTAL", f"=B44+B45+B46", money_fmt),
    ("Savings vs today", f"=B31-B47", money_fmt),
    ("% reduction", f"=1-B47/B31", pct_fmt),
    ("vs Trigger.dev", f'=IF(B47<=B56,"Render cheaper by "&TEXT(1-B47/B56,"0%"),"Trigger cheaper by "&TEXT(1-B56/B47,"0%"))', "@"),
]
for i, (label, formula, fmt) in enumerate(approach2):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    if label == "TOTAL":
        style_result(ws3, row, 2).value = formula
    else:
        style_cell(ws3, row, 2).value = formula
    ws3.cell(row, 2).number_format = fmt

# ── TRIGGER.DEV ──
r = 52
ws3.cell(r, 1, "TRIGGER.DEV").font = section_font
ws3.cell(r, 1).fill = PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid")
r = 53
trigger = [
    ("Parent compute cost (active only)", f"={req_cell}*{p_ac_cell}*{t_rate}", money_fmt),
    ("Subtask compute cost (active only)", f"={req_cell}*{subs_cell}*{s_ac_cell}*{t_rate}", money_fmt),
    ("Invocation fees", f"={req_cell}*(1+{subs_cell})*{t_invoc}", money_fmt),
    ("TOTAL", f"=B53+B54+B55", money_fmt),
]
for i, (label, formula, fmt) in enumerate(trigger):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    if label == "TOTAL":
        style_result(ws3, row, 2).value = formula
    else:
        style_cell(ws3, row, 2).value = formula
    ws3.cell(row, 2).number_format = fmt

# ── BG WORKERS ──
r = 58
ws3.cell(r, 1, "BACKGROUND WORKERS").font = section_font
r = 59
bg = [
    ("Instance type cost ($/month)", 175, money_fmt),
    ("Number of instances", 4, int_fmt),
    ("TOTAL", f"=B59*B60", money_fmt),
]
for i, (label, val, fmt) in enumerate(bg):
    row = r + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    if label == "TOTAL":
        style_result(ws3, row, 2).value = val
    elif isinstance(val, str):
        style_cell(ws3, row, 2).value = val
    else:
        style_input(ws3, row, 2).value = val
    ws3.cell(row, 2).number_format = fmt

# ── SUMMARY ──
r = 63
ws3.cell(r, 1, "SUMMARY COMPARISON").font = Font(bold=True, size=14, color="2F5496")
r = 64
for c, h in enumerate(["Option", "Monthly Cost", "vs Today", "vs Trigger.dev"], 1):
    ws3.cell(r, c, h)
style_header(ws3, r, 4)

summary_rows = [
    ("Render Workflows (today)", "=B31", "", ""),
    ("Approach 1 (no parent billing)", "=B38", "=B31-B38", "=B38-B56"),
    ("Approach 2 (active compute)", "=B47", "=B31-B47", "=B47-B56"),
    ("Approach 3 (waitpoints on top of 2)", "=B47", "", ""),
    ("  ^ same baseline, further savings", "", "", ""),
    ("  ^ on long LLM/API waits + new", "", "", ""),
    ("  ^ use cases (approval flows, etc)", "", "", ""),
    ("Trigger.dev", "=B56", "", ""),
    ("Background Workers", "=B61", "", ""),
]
for i, (label, cost, vs_today, vs_trigger) in enumerate(summary_rows):
    row = r + 1 + i
    style_cell(ws3, row, 1).value = label
    ws3.cell(row, 1).font = label_font
    style_result(ws3, row, 2).value = cost
    ws3.cell(row, 2).number_format = money_fmt
    if vs_today:
        style_cell(ws3, row, 3).value = vs_today
        ws3.cell(row, 3).number_format = money_fmt
    if vs_trigger:
        style_cell(ws3, row, 4).value = vs_trigger
        ws3.cell(row, 4).number_format = money_fmt

ws3.column_dimensions["A"].width = 42
ws3.column_dimensions["B"].width = 20
ws3.column_dimensions["C"].width = 16
ws3.column_dimensions["D"].width = 16


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 4: APPROACH COMPARISON
# ═══════════════════════════════════════════════════════════════════════════

ws4 = wb.create_sheet("Approaches")
ws4.sheet_properties.tabColor = "C00000"

ws4["A1"] = "Three Approaches — Tradeoff Comparison"
ws4["A1"].font = Font(bold=True, size=14, color="2F5496")
ws4.merge_cells("A1:D1")

r = 3
headers = ["Dimension", "Approach 1\nNo parent billing", "Approach 2\nActive compute only", "Approach 3\nWaitpoints (checkpoint)"]
for c, h in enumerate(headers, 1):
    ws4.cell(r, c, h)
style_header(ws4, r, 4)

rows = [
    ("What changes", "Billing only", "Billing + runtime metering", "SDK + checkpointing infra (on top of Approach 2)"),
    ("Customer cost (AI case)", "=\'Customer Scenario\'!B38", "=\'Customer Scenario\'!B47", "Approach 2 baseline + further savings on long LLM waits"),
    ("% reduction vs today", "=\'Customer Scenario\'!B40", "=\'Customer Scenario\'!B49", "86%+ (depends on wait duration)"),
    ("vs Trigger.dev", "Still more expensive", "Cheaper (lower rate, no invoc fee)", "Even cheaper for long-running calls"),
    ("SDK changes required", "None", "None", "Yes — new wait primitives (wait.for_result, wait.for_callback, wait.for_duration)"),
    ("Customer code changes", "None", "None", "Wrap long LLM/API calls in wait hooks"),
    ("Performance impact", "None", "None", "Adds checkpoint/restore latency per waitpoint"),
    ("Implementation complexity", "Low", "Medium", "High — requires Approach 2 first"),
    ("Key dependency", "Billing system", "cgroup CPU tracking", "CRIU or equivalent + checkpoint storage"),
    ("Time to ship (estimate)", "Weeks", "1-2 months", "3-6 months (after Approach 2)"),
    ("Revenue impact", "Moderate (~30-50%)", "Large (~60-80%)", "Incremental beyond Approach 2"),
    ("New use cases unlocked", "No", "No", "Yes — approval flows, webhook waits, multi-day workflows, long-running agents"),
    ("Closes pricing gap alone?", "No — subtasks still bill wall-clock", "Yes — fully competitive", "No — builds on Approach 2, optimizes further for long waits"),
    ("Relationship to other approaches", "Standalone", "Standalone (supercedes Approach 1)", "Additive — requires Approach 2 as foundation"),
]

for i, (dim, a1, a2, a3) in enumerate(rows):
    row = r + 1 + i
    style_cell(ws4, row, 1).value = dim
    ws4.cell(row, 1).font = label_font
    for c, val in enumerate([a1, a2, a3], 2):
        cell = style_cell(ws4, row, c)
        cell.value = val
        cell.alignment = Alignment(wrap_text=True)
        if isinstance(val, str) and val.startswith("="):
            if "B40" in val or "B49" in val:
                cell.number_format = pct_fmt
            else:
                cell.number_format = money_fmt

ws4.column_dimensions["A"].width = 28
ws4.column_dimensions["B"].width = 28
ws4.column_dimensions["C"].width = 28
ws4.column_dimensions["D"].width = 32


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 5: RENDER WINS — Which workloads Render wins by project
# ═══════════════════════════════════════════════════════════════════════════

ws5 = wb.create_sheet("Render Wins")
ws5.sheet_properties.tabColor = "548235"

ws5["A1"] = "When Does Workflows Win?"
ws5["A1"].font = Font(bold=True, size=14, color="2F5496")
ws5.merge_cells("A1:L1")
ws5["A2"] = "Three projects, progressive impact. Yellow = inputs you can change."
ws5["A2"].font = Font(italic=True, size=9, color="666666")

# ── INPUTS ──
r = 4
ws5.cell(r, 1, "RATES").font = section_font
r = 5
rate_labels = [
    ("Workflows $/second", 0.0000556, rate_fmt, "(Standard: 1 CPU, 2 GB)"),
    ("Trigger.dev $/second", 0.0000850, rate_fmt, "(Medium 1x: 1 CPU, 2 GB — same specs)"),
    ("Trigger.dev invocation $/run", 0.000025, rate_fmt, "(charged per task run, Workflows has none)"),
    ("BG Worker $/month per instance", 85, money_fmt, "(Pro: 2 CPU, 4 GB)"),
    ("BG subprocess workers per instance", 2, int_fmt, "(e.g., Gunicorn/Celery workers — 1 per CPU, ~1GB RAM each)"),
    ("BG async slots per subprocess", 50, int_fmt, "(concurrent async I/O requests per subprocess)"),
    ("Peak-to-average traffic ratio", 3, '0.0"x"', "(traffic skews to business hours)"),
]
for i, (label, val, fmt, note) in enumerate(rate_labels):
    row = r + i
    style_cell(ws5, row, 1).value = label
    ws5.cell(row, 1).font = label_font
    style_input(ws5, row, 2).value = val
    ws5.cell(row, 2).number_format = fmt
    ws5.cell(row, 3, note).font = Font(size=9, italic=True, color="888888")

rw_rate = "$B$5"   # Workflows rate
tw_rate = "$B$6"   # Trigger rate
tw_invoc = "$B$7"  # Trigger invocation
bg_mo = "$B$8"     # BG Worker monthly
bg_subprocs = "$B$9"  # subprocess workers per instance
bg_slots = "$B$10"  # async slots per subprocess
bg_peak = "$B$11"  # peak-to-average ratio

# ── MAIN TABLE ──
r = 13
ws5.cell(r, 1, "COST COMPARISON BY WORKLOAD").font = section_font

r = 14
headers = [
    "Task Type", "Typically\nRuns On",
    "Monthly\nVolume", "Wall-Clock\n(s)", "Active\nCompute (s)", "Waitpoint\nCompute (s)", "Runs/\nReq",
]
for c, h in enumerate(headers, 1):
    ws5.cell(r, c, h)
style_header(ws5, r, 7)

# Cost columns header (separate color)
cost_headers = [
    "Workflows\nTODAY",
    "P1: No Parent\nBilling",
    "P2: Active\nCompute Only",
    "P3: Waitpoints",
    "Trigger.dev",
    "BG Instances\nNeeded", "BG Workers\nCost",
    "Project Needed\nto Win"
]
for c, h in enumerate(cost_headers, 8):
    ws5.cell(r, c, h)
    ws5.cell(r, c).font = header_font
    ws5.cell(r, c).fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
    ws5.cell(r, c).alignment = Alignment(horizontal="center", wrap_text=True)
    ws5.cell(r, c).border = thin_border

tasks = [
    ("Simple compute",        "Workflows",   500000,  1,    1,    1,    1),
    ("Image processing",      "BG Worker",   200000,  5,    5,    5,    1),
    ("Single LLM call",       "BG Worker",   300000,  8,    0.5,  0.3,  1),
    ("Multi-step LLM agent",  "BG Worker",   100000,  35,   2.5,  1.5,  6),
    ("AI pipeline",           "BG Worker",    50000,  90,   6,    3,    4),
    ("Data fan-out",          "Workflows",    10000,  200,  90,   85,   121),
    ("Long-running agent",    "Trigger.dev",  10000,  350,  15,   5,    10),
]

data_start = 15
for i, (name, typical, volume, wc, ac, wp, runs) in enumerate(tasks):
    row = data_start + i

    style_cell(ws5, row, 1).value = name
    ws5.cell(row, 1).font = label_font

    # Typically runs on (color-coded)
    style_cell(ws5, row, 2).value = typical
    color_map = {"BG Worker": "FFEB9C", "Trigger.dev": "E8DAEF", "Workflows": "DAEEF3"}
    ws5.cell(row, 2).fill = PatternFill(start_color=color_map.get(typical, "FFFFFF"),
                                         end_color=color_map.get(typical, "FFFFFF"), fill_type="solid")

    style_input(ws5, row, 3).value = volume
    ws5.cell(row, 3).number_format = int_fmt
    style_input(ws5, row, 4).value = wc
    style_input(ws5, row, 5).value = ac
    style_input(ws5, row, 6).value = wp
    style_input(ws5, row, 7).value = runs

    vol = f"C{row}"

    # Col H: Workflows TODAY = vol × wall_clock × rate
    style_result(ws5, row, 8)
    ws5.cell(row, 8).value = f"={vol}*D{row}*{rw_rate}"
    ws5.cell(row, 8).number_format = money_fmt

    # Col I: P1 — parent billed active, subtasks still wall-clock
    style_result(ws5, row, 9)
    ws5.cell(row, 9).value = (
        f"=IF(G{row}=1,H{row},"
        f"{vol}*(E{row}+(D{row}-E{row})*(G{row}-1)/G{row})*{rw_rate})"
    )
    ws5.cell(row, 9).number_format = money_fmt

    # Col J: P2 — all tasks billed active compute only
    style_result(ws5, row, 10)
    ws5.cell(row, 10).value = f"={vol}*E{row}*{rw_rate}"
    ws5.cell(row, 10).number_format = money_fmt

    # Col K: P3 — waitpoint compute (even less than active)
    style_result(ws5, row, 11)
    ws5.cell(row, 11).value = f"={vol}*F{row}*{rw_rate}"
    ws5.cell(row, 11).number_format = money_fmt

    # Col L: Trigger.dev (already has active compute + waitpoints)
    style_result(ws5, row, 12)
    ws5.cell(row, 12).value = f"={vol}*F{row}*{tw_rate}+{vol}*G{row}*{tw_invoc}"
    ws5.cell(row, 12).number_format = money_fmt

    # Col M: BG Instances needed = MAX(CPU-bound, concurrency-bound)
    # CPU-bound: total CPU-seconds / (730h × 3600s × subprocesses_per_instance)
    # Concurrency: peak_concurrent_requests / (subprocesses × async_slots)
    style_result(ws5, row, 13)
    ws5.cell(row, 13).value = (
        f"=MAX("
        f"CEILING({vol}*E{row}/(730*3600*{bg_subprocs}),1),"
        f"CEILING(({vol}/730/3600)*D{row}*{bg_peak}/({bg_subprocs}*{bg_slots}),1)"
        f")"
    )
    ws5.cell(row, 13).number_format = int_fmt

    # Col N: BG Workers cost = instances × monthly
    style_result(ws5, row, 14)
    ws5.cell(row, 14).value = f"=M{row}*{bg_mo}"
    ws5.cell(row, 14).number_format = money_fmt

    # Col O: Which project makes Workflows cheapest?
    style_result(ws5, row, 15)
    ws5.cell(row, 15).value = (
        f'=IF(AND(H{row}<=L{row},H{row}<=N{row}),"Already wins",'
        f'IF(AND(I{row}<=L{row},I{row}<=N{row}),"P1: No parent billing",'
        f'IF(AND(J{row}<=L{row},J{row}<=N{row}),"P2: Active compute",'
        f'IF(AND(K{row}<=L{row},K{row}<=N{row}),"P3: Waitpoints",'
        f'"Cannot win"))))'
    )
    ws5.cell(row, 15).font = Font(bold=True, size=10)

data_end = data_start + len(tasks) - 1

# ── SCORECARD ──
r = data_end + 2
ws5.cell(r, 1, "SCORECARD").font = Font(bold=True, size=12, color="2F5496")
r += 1
score_headers = ["", "Workflows", "Trigger.dev", "BG Worker"]
for c, h in enumerate(score_headers, 1):
    ws5.cell(r, c, h)
style_header(ws5, r, 4)

score_rows = [
    ("Cheapest today",
     f'=COUNTIFS(O{data_start}:O{data_end},"Already wins")',
     f'=COUNTIFS(N{data_start}:N{data_end},"<>Already wins",'
     f'N{data_start}:N{data_end},"<>P1*",'
     f'N{data_start}:N{data_end},"<>P2*",'
     f'N{data_start}:N{data_end},"<>P3*",'
     f'N{data_start}:N{data_end},"<>Cannot win")',
     None),  # complex, use simpler below
    ("After P1: No parent billing", None, None, None),
    ("After P2: Active compute only", None, None, None),
    ("After P3: Waitpoints", None, None, None),
]

# Simpler scorecard: count wins at each stage (col O = "Project Needed to Win")
oc = f"O{data_start}:O{data_end}"
r += 1
ws5.cell(r, 1, "Workflows wins (today)").font = label_font
style_result(ws5, r, 2).value = f'=COUNTIF({oc},"Already wins")'
ws5.cell(r, 2).number_format = '0" of 7"'

r += 1
ws5.cell(r, 1, "Workflows wins (after P1)").font = label_font
style_result(ws5, r, 2).value = f'=COUNTIF({oc},"Already wins")+COUNTIF({oc},"P1*")'
ws5.cell(r, 2).number_format = '0" of 7"'

r += 1
ws5.cell(r, 1, "Workflows wins (after P1+P2)").font = label_font
style_result(ws5, r, 2).value = f'=COUNTIF({oc},"Already wins")+COUNTIF({oc},"P1*")+COUNTIF({oc},"P2*")'
ws5.cell(r, 2).number_format = '0" of 7"'

r += 1
ws5.cell(r, 1, "Workflows wins (after P1+P2+P3)").font = label_font
style_result(ws5, r, 2).value = f'=COUNTIF({oc},"Already wins")+COUNTIF({oc},"P1*")+COUNTIF({oc},"P2*")+COUNTIF({oc},"P3*")'
ws5.cell(r, 2).number_format = '0" of 7"'

r += 1
ws5.cell(r, 1, "Cannot win").font = label_font
style_result(ws5, r, 2).value = f'=COUNTIF({oc},"Cannot win")'
ws5.cell(r, 2).number_format = '0" of 7"'
ws5.cell(r, 2).fill = warning_fill

# ── NOTES ──
r += 2
ws5.cell(r, 1, "DEFINITIONS").font = section_font
r += 1
notes = [
    "P1 — No parent billing: Stop billing parent tasks while subtasks run. Billing change only. Ships in weeks.",
    "P2 — Active compute only: Bill CPU-active time only across all tasks. Needs cgroup tracking. Ships in 1-2 months.",
    "P3 — Waitpoints: Explicit SDK primitives (wait.for_result, wait.for_callback) that release containers.",
    "     Needs checkpointing infra (CRIU). Ships in 3-6 months. Builds on P2.",
    "",
    "Trigger.dev already has P2 + P3 (active compute billing with CRIU checkpointing).",
    "Their cost column uses 'Waitpoint Compute' because their billing reflects both optimizations.",
    "",
    "BG Worker is a flat monthly cost regardless of workload. Wins when per-use costs exceed the fixed rate.",
    "Does NOT include managed retries, orchestration, observability, or burst handling.",
    "'Project Needed to Win' shows the FIRST project that makes Workflows the cheapest of all three options.",
]
for i, note in enumerate(notes):
    ws5.cell(r + i, 1, note).font = Font(size=9, italic=True, color="666666")
    ws5.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=10)

for c in range(1, 16):
    widths = [24, 12, 12, 10, 10, 10, 7, 14, 14, 14, 14, 14, 10, 12, 24]
    ws5.column_dimensions[get_column_letter(c)].width = widths[c - 1]


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 6: WHO WINS — Clean view of winner at each stage
# ═══════════════════════════════════════════════════════════════════════════

ws6 = wb.create_sheet("Who Wins")
ws6.sheet_properties.tabColor = "BF8F00"

ws6["A1"] = "Who Is Cheapest? — By Workload Type and Development Stage"
ws6["A1"].font = Font(bold=True, size=14, color="2F5496")
ws6.merge_cells("A1:H1")

# ── ASSUMPTIONS ──
r = 3
ws6.cell(r, 1, "ASSUMPTIONS").font = section_font

assumptions_data = [
    ("Workflows instance", "Standard (1 CPU, 2 GB) at $0.20/hr = $0.0000556/s"),
    ("Trigger.dev instance", "Medium 1x (1 CPU, 2 GB) at $0.0000850/s + $0.000025/run invocation"),
    ("BG Worker instance", "Pro (2 CPU, 4 GB) at $85/month, 1 instance"),
    ("BG Worker capacity", "2 subprocess workers × 50 async slots = 100 concurrent requests"),
    ("Traffic pattern", "3× peak-to-average ratio (business hours skew)"),
    ("Trigger.dev billing", "Already has active compute billing AND waitpoints (CRIU checkpointing)"),
    ("Workflows billing (today)", "Bills full wall-clock time including waits for subtasks and API calls"),
    ("BG Worker features", "NO managed retries, orchestration, observability, or burst handling"),
]
for i, (label, desc) in enumerate(assumptions_data):
    row = r + 1 + i
    ws6.cell(row, 1, label).font = Font(bold=True, size=9)
    ws6.cell(row, 2, desc).font = Font(size=9, color="444444")
    ws6.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)

# ── WORKLOAD INPUTS ──
r = r + len(assumptions_data) + 2
ws6.cell(r, 1, "WORKLOAD PROFILES (yellow = editable)").font = section_font
r += 1

input_headers = ["Task Type", "Monthly Volume", "Wall-Clock (s)", "Active Compute (s)",
                  "Waitpoint Compute (s)", "Task Runs / Request", "Wait Ratio"]
for c, h in enumerate(input_headers, 1):
    ws6.cell(r, c, h)
style_header(ws6, r, len(input_headers))

tasks = [
    ("Simple compute",        500000,  1,    1,    1,    1),
    ("Image processing",      200000,  5,    5,    5,    1),
    ("Single LLM call",       300000,  8,    0.5,  0.3,  1),
    ("Multi-step LLM agent",  100000,  35,   2.5,  1.5,  6),
    ("AI pipeline",            50000,  90,   6,    3,    4),
    ("Data fan-out",           10000,  200,  90,   85,   121),
    ("Long-running agent",     10000,  350,  15,   5,    10),
]

input_start = r + 1
for i, (name, vol, wc, ac, wp, runs) in enumerate(tasks):
    row = input_start + i
    style_cell(ws6, row, 1).value = name
    ws6.cell(row, 1).font = label_font
    style_input(ws6, row, 2).value = vol
    ws6.cell(row, 2).number_format = int_fmt
    style_input(ws6, row, 3).value = wc
    style_input(ws6, row, 4).value = ac
    style_input(ws6, row, 5).value = wp
    style_input(ws6, row, 6).value = runs
    # Wait ratio = 1 - active/wallclock
    style_cell(ws6, row, 7).value = f"=1-D{row}/C{row}"
    ws6.cell(row, 7).number_format = pct_fmt

input_end = input_start + len(tasks) - 1

# ── WHO WINS TABLE ──
r = input_end + 2
ws6.cell(r, 1, "WHO IS CHEAPEST AT EACH STAGE").font = section_font
ws6.cell(r + 1, 1, "Each cell shows which product is cheapest for that workload at that development stage.").font = Font(size=9, italic=True, color="666666")
ws6.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=6)

r += 2
win_headers = ["Task Type", "Today\n(wall-clock billing)",
               "After P1\nNo parent billing",
               "After P2\nActive compute only",
               "After P3\nWaitpoints",
               "What Changes\nat Winning Stage"]
for c, h in enumerate(win_headers, 1):
    ws6.cell(r, c, h)
style_header(ws6, r, len(win_headers))

# Rates for formulas (hardcoded to match assumptions above)
RR = 0.0000556
TR = 0.0000850
TI = 0.000025
BG = 85

win_start = r + 1
for i, (name, vol, wc, ac, wp, runs) in enumerate(tasks):
    row = win_start + i
    irow = input_start + i  # corresponding input row

    style_cell(ws6, row, 1).value = name
    ws6.cell(row, 1).font = label_font

    # Calculate costs for formulas — reference input rows
    # Workflows today
    wf_today = f"B{irow}*C{irow}*0.0000556"
    # P1: parent active, subtasks wall-clock
    wf_p1 = f"IF(F{irow}=1,{wf_today},B{irow}*(D{irow}+(C{irow}-D{irow})*(F{irow}-1)/F{irow})*0.0000556)"
    # P2: all active compute
    wf_p2 = f"B{irow}*D{irow}*0.0000556"
    # P3: waitpoint compute
    wf_p3 = f"B{irow}*E{irow}*0.0000556"
    # Trigger (has active compute + waitpoints)
    trig = f"B{irow}*E{irow}*0.0000850+B{irow}*F{irow}*0.000025"
    # BG Worker
    bgw = "85"

    # Col B: Winner today
    style_result(ws6, row, 2)
    ws6.cell(row, 2).value = (
        f'=IF(AND({wf_today}<={trig},{wf_today}<={bgw}),"Workflows",'
        f'IF(AND({trig}<={wf_today},{trig}<={bgw}),"Trigger.dev","BG Worker"))'
    )

    # Col C: Winner after P1
    style_result(ws6, row, 3)
    ws6.cell(row, 3).value = (
        f'=IF(AND({wf_p1}<={trig},{wf_p1}<={bgw}),"Workflows",'
        f'IF(AND({trig}<={wf_p1},{trig}<={bgw}),"Trigger.dev","BG Worker"))'
    )

    # Col D: Winner after P2
    style_result(ws6, row, 4)
    ws6.cell(row, 4).value = (
        f'=IF(AND({wf_p2}<={trig},{wf_p2}<={bgw}),"Workflows",'
        f'IF(AND({trig}<={wf_p2},{trig}<={bgw}),"Trigger.dev","BG Worker"))'
    )

    # Col E: Winner after P3
    style_result(ws6, row, 5)
    ws6.cell(row, 5).value = (
        f'=IF(AND({wf_p3}<={trig},{wf_p3}<={bgw}),"Workflows",'
        f'IF(AND({trig}<={wf_p3},{trig}<={bgw}),"Trigger.dev","BG Worker"))'
    )

    # Col F: What changes — describe the transition
    style_result(ws6, row, 6)
    ws6.cell(row, 6).value = (
        f'=IF(B{row}="Workflows","Already cheapest — no change needed",'
        f'IF(AND(B{row}<>"Workflows",C{row}="Workflows"),"P1 stops parent wait billing → flips from "&B{row}'
        f',IF(AND(C{row}<>"Workflows",D{row}="Workflows"),"P2 bills active compute only → flips from "&C{row}'
        f',IF(AND(D{row}<>"Workflows",E{row}="Workflows"),"P3 waitpoints pause LLM billing → flips from "&D{row}'
        f',"Workflows does not win even after P3"))))'
    )
    ws6.cell(row, 6).alignment = Alignment(wrap_text=True)

win_end = win_start + len(tasks) - 1

# ── SCORECARD ──
r = win_end + 2
ws6.cell(r, 1, "SCORECARD — Workflows wins by stage").font = Font(bold=True, size=12, color="2F5496")
r += 1

sc_headers = ["Stage", "Workflows Wins", "Trigger.dev Wins", "BG Worker Wins"]
for c, h in enumerate(sc_headers, 1):
    ws6.cell(r, c, h)
style_header(ws6, r, 4)

stages = [
    ("Today", "B"),
    ("After P1: No parent billing", "C"),
    ("After P2: Active compute only", "D"),
    ("After P3: Waitpoints", "E"),
]
for i, (label, col) in enumerate(stages):
    row = r + 1 + i
    ws6.cell(row, 1, label).font = label_font
    rng = f"{col}{win_start}:{col}{win_end}"
    style_result(ws6, row, 2).value = f'=COUNTIF({rng},"Workflows")'
    ws6.cell(row, 2).number_format = '0" of 7"'
    style_result(ws6, row, 3).value = f'=COUNTIF({rng},"Trigger.dev")'
    ws6.cell(row, 3).number_format = '0" of 7"'
    style_result(ws6, row, 4).value = f'=COUNTIF({rng},"BG Worker")'
    ws6.cell(row, 4).number_format = '0" of 7"'

# ── PROJECT DESCRIPTIONS ──
r = r + len(stages) + 3
ws6.cell(r, 1, "PROJECT DEFINITIONS").font = section_font
r += 1
proj_notes = [
    ("P1 — No parent billing", "Stop billing parent tasks while awaiting subtasks. Billing change only. No SDK, runtime, or customer code changes. Ships in weeks."),
    ("P2 — Active compute only", "Bill CPU-active time across ALL tasks (parent + subtasks). Requires cgroup CPU tracking in runtime. No customer code changes. Ships in 1-2 months."),
    ("P3 — Waitpoints", "Explicit SDK primitives (wait.for_result, wait.for_callback) that release containers during long LLM/API calls. Requires CRIU checkpointing infrastructure. Customer wraps API calls in wait hooks. Ships in 3-6 months. Builds on P2."),
]
for i, (title, desc) in enumerate(proj_notes):
    row = r + i
    ws6.cell(row, 1, title).font = Font(bold=True, size=9)
    ws6.cell(row, 2, desc).font = Font(size=9, color="444444")
    ws6.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)

for c in range(1, 7):
    ws6.column_dimensions[get_column_letter(c)].width = [24, 18, 18, 18, 18, 40][c - 1]


# ── Save ────────────────────────────────────────────────────────────────────

out = "/Users/yaroslavborets/Dev/WorkflowsDemo/workflow-demo-test-web/COST_MODEL.xlsx"
wb.save(out)
print(f"Saved to {out}")
